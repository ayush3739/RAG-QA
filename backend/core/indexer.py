import asyncio
import hashlib
import os
import pickle
from pathlib import Path

from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from sqlalchemy.ext.asyncio import AsyncSession
from langchain_core.documents import Document
import docx
from sqlalchemy import select
from backend.models import models


from backend.core.config import settings
from backend.core.utils import simple_tokenize
from backend.models.models import Chunk
from rank_bm25 import BM25Okapi


class Indexer:
    def __init__(self, file_path: str, db_session: AsyncSession, document_id: int,document_public_id :int):
        self.file_path = Path(file_path)
        self.db = db_session
        self.document_id = document_id
        self.document_public_id = document_public_id
        self.embedding_model = OpenAIEmbeddings(
            api_key=settings.github_token,
            model="text-embedding-3-small",
            openai_api_base="https://models.github.ai/inference",
        )

    async def index(self):
        try:
            ext = self.file_path.suffix.lower()
            if ext == ".pdf":
                # Step 1: load PDF — CPU/IO bound, offload to thread
                loader = PyPDFLoader(file_path=str(self.file_path))
                docs = await asyncio.to_thread(loader.load)
                for d in docs:
                    d.metadata["source"] = self.file_path.name
            elif ext == ".docx":
                # load Word DOCX to text
                text = await asyncio.to_thread(self._read_docx, self.file_path)
                docs = [Document(page_content=text, metadata={"source": self.file_path.name})]
            elif ext in {".txt", ".md"}:
                # load Plain Text or Markdown with default page 1
                text = await asyncio.to_thread(self._read_plain_text, self.file_path)
                docs = [Document(
                    page_content=text,
                    metadata={
                        "source": self.file_path.name,
                        "page": 1,
                        "page_label": "1"
                    }
                )]
            elif ext == ".doc":
                raise ValueError("Older binary .doc format is not supported for indexing. Please convert to .docx.")
            else:
                raise ValueError(f"Unsupported file format: {ext}")

            # Step 2: chunk — CPU bound, offload to thread
            text_splitter = RecursiveCharacterTextSplitter(
                chunk_size=600,
                chunk_overlap=150,
            )
            chunks = await asyncio.to_thread(text_splitter.split_documents, docs)
            
            # Guard against huge files to avoid huge billing or OOM
            if len(chunks) > 10000:
                raise ValueError("Document too large (exceeded 10,000 chunks limit).")
            print(f"Total chunks created: {len(chunks)}")

            # Stamp deterministic chunk_id on each chunk
            for chunk in chunks:
                normalized_text = " ".join(chunk.page_content.split())
                page_label = chunk.metadata.get("page_label", "")
                source = chunk.metadata.get("source", self.file_path.name)
                # Ensure the metadata has source set to filename
                chunk.metadata["source"] = source
                chunk.metadata["chunk_id"] = self.make_chunk_id(
                    normalized_text, page_label, source
                )

            # Step 3: BM25 — CPU bound, offload to thread
            await asyncio.to_thread(self._persist_bm25, chunks)

            # Step 4: embed — Langchain OpenAIEmbeddings batches automatically internally
            texts = [c.page_content for c in chunks]
            embeddings = await asyncio.to_thread(
                self.embedding_model.embed_documents,
                texts,
            )

            # Step 5: bulk insert chunks — pure async DB write
            rows = [
                Chunk(
                    document_id=self.document_id,
                    chunk_id=ch.metadata["chunk_id"],
                    page_content=ch.page_content,
                    page_number=ch.metadata.get("page"),
                    chunk_index=idx,
                    source=ch.metadata.get("source"),
                    embedding=emb,
                )
                for idx, (ch, emb) in enumerate(zip(chunks, embeddings))
            ]

            self.db.add_all(rows)
            await self.db.flush()
            
            # Update parent Document's status and count within the transaction
            result = await self.db.execute(
                select(models.Document).where(models.Document.id == self.document_id)
            )
            doc = result.scalars().first()
            if doc:
                doc.chunk_count = len(rows)
                doc.status = "indexed"
                
            await self.db.commit()
            print(f"✓ {len(rows)} chunks committed to DB for document_id={self.document_id}")

        except Exception as e:
            await self.db.rollback()
            raise RuntimeError(f"Indexing failed: {e}")

    def _persist_bm25(self, chunks) -> None:
        """Sync — called via asyncio.to_thread."""
        try:
            texts = [c.page_content for c in chunks]
            tokenized = [simple_tokenize(t) for t in texts]
            bm25 = BM25Okapi(tokenized)
            meta = [
                {
                    "chunk_id": c.metadata.get("chunk_id"),
                    "page_content": t,
                    "page_label": c.metadata.get("page_label"),
                    "source": c.metadata.get("source"),
                }
                for t, c in zip(texts, chunks)
            ]
            os.makedirs("data/bm25", exist_ok=True)
            path = f"data/bm25/{self.document_public_id}_bm25.pkl"
            with open(path, "wb") as f:
                pickle.dump({"bm25": bm25, "meta": meta}, f)
            print(f"✓ BM25 persisted → {path}")
        except Exception as e:
            print(f"Warning: BM25 persist failed: {e}")

    def make_chunk_id(self, text: str, page_label: str, source: str) -> str:
        h = hashlib.sha256(
            f"{source}|{page_label}|{text}".encode("utf-8")
        ).hexdigest()
        return h[:16]

    def _read_docx(self, path: Path) -> str:
        try:
            doc = docx.Document(path)
            paragraphs_and_tables = []
            
            for element in doc.element.body:
                if element.tag.endswith('p'):
                    p = docx.text.paragraph.Paragraph(element, doc)
                    if p.text.strip():
                        paragraphs_and_tables.append(p.text)
                elif element.tag.endswith('tbl'):
                    t = docx.table.Table(element, doc)
                    table_text = []
                    for row in t.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            table_text.append(" | ".join(row_text))
                    if table_text:
                        paragraphs_and_tables.append("\n".join(table_text))
                        
            return "\n\n".join(paragraphs_and_tables)
        except Exception as e:
            raise ValueError(f"Failed to parse Word Document (.docx): {e}")

    def _read_plain_text(self, path: Path) -> str:
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            raise ValueError(f"Failed to read text file: {e}")